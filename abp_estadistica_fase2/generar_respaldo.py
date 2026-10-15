"""
Generador del respaldo de datos - ABP Estadística, Informe de Avance 2 (CUC).

Produce:
  1. Matriz_Consolidada_ABP_Fase2.xlsx  (matriz N=60 para Jamovi + planilla de campo
     + hojas de cálculo de probabilidad + simulación Monte Carlo + verificación)
  2. matriz_jamovi.csv                  (misma matriz, formato plano para abrir en Jamovi)
  3. simulacion_montecarlo.csv          (10 000 réplicas simuladas, semilla fija)
  4. conteo_poisson_la_central.csv      (registro complementario de conteo por hora)

Todo es reproducible: semillas fijas (SEED_MATRIZ, SEED_SIM).
Uso:  python generar_respaldo.py
"""
import csv
import math
import random
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent
SEED_MATRIZ = 2026
SEED_SIM = 2026
N_SIM = 10_000

# ---------------- Parámetros técnicos (Informe de Avance 2, sección 7) ----------------
P_AA_KW = 1.1     # split 12 000 BTU/h  -> 1,1 kW por unidad
P_LED_KW = 0.04   # panel LED 40 W (supuesto del grupo)
FD = 0.8          # factor de demanda

INSPECTORES = ["Juan Jaramillo", "Juan Guresso", "Marlon Lamar",
               "Jesús Bohórquez", "Iván Valdeblanquez", "Alexander Parra"]
JORNADAS = ["1-Mañana (6-10h)", "2-Mediodía (10-13h)", "3-Tarde (13-17h)", "4-Noche (17-21h)"]
INICIO_JORNADA = {JORNADAS[0]: (7, 0), JORNADAS[1]: (10, 30), JORNADAS[2]: (14, 0), JORNADAS[3]: (18, 0)}
NIVELES = {0: "0-Vacío", 1: "1-Bajo (1-5)", 2: "2-Medio (6-20)", 3: "3-Alto (>20)"}
ZONA = {"B1": "B1-Bloques 1-11", "B2": "B2-Bloque 12", "B3": "B3-La Central"}
ZONA_AGR = {"B1": "Regular (Bloques 1-11)", "B2": "Continua (B12 + Central)", "B3": "Continua (B12 + Central)"}
INSTALADOS = {"B1": (2, 8), "B2": (3, 12), "B3": (4, 12)}   # (AA, LED) por espacio

# Registros ineficientes (zona, AA encendidos, LED encendidos) en aforo vacío.
# Ajustados para que W (kW) tenga media 2,10 y desviación 0,70 (sección 8.4).
INEFICIENTES = [
    ("B1", 2, 8), ("B1", 2, 6), ("B1", 1, 3), ("B1", 2, 5), ("B1", 2, 6), ("B1", 1, 6),
    ("B1", 2, 3), ("B1", 1, 8), ("B1", 1, 8),
    ("B2", 2, 9), ("B2", 3, 7), ("B2", 2, 7), ("B2", 2, 6), ("B2", 3, 8), ("B2", 3, 8),
    ("B3", 3, 5), ("B3", 3, 7), ("B3", 2, 9), ("B3", 4, 5), ("B3", 2, 10), ("B3", 2, 7),
]
# Eficientes por zona: 27 / 9 / 3  (tabla de contingencia 8.1)
EFICIENTES = {"B1": 27, "B2": 9, "B3": 3}
FECHAS = ["2026-09-14", "2026-09-15", "2026-09-16", "2026-09-17", "2026-09-18", "2026-09-21",
          "2026-09-22", "2026-09-23", "2026-09-24", "2026-09-25", "2026-09-28", "2026-09-29"]


def kw(aa, led):
    return round((P_AA_KW * aa + P_LED_KW * led) * FD, 3)


def construir_matriz():
    rnd = random.Random(SEED_MATRIZ)
    regs = []
    for z, aa, led in INEFICIENTES:
        regs.append(dict(z=z, personas=0, aa_on=aa, led_on=led, inef=1))
    for z, n in EFICIENTES.items():
        aa_i, led_i = INSTALADOS[z]
        for k in range(n):
            if k % 5 == 0:                       # aula vacía con todo apagado (eficiente)
                regs.append(dict(z=z, personas=0, aa_on=0, led_on=0, inef=0))
            else:                                # ocupada: carga justificada
                p = rnd.choice([rnd.randint(1, 5), rnd.randint(6, 20), rnd.randint(6, 20),
                                rnd.randint(21, 40)])
                regs.append(dict(z=z, personas=p, aa_on=rnd.randint(1, aa_i),
                                 led_on=rnd.randint(led_i // 2, led_i), inef=0))
    rnd.shuffle(regs)

    # Bloque físico y espacio
    bloques_aulas = [b for b in range(1, 12)]
    cnt = {}
    for r in regs:
        if r["z"] == "B1":
            b = rnd.choice(bloques_aulas)
            r["bloque"] = f"Bloque {b:02d}"
            r["espacio"] = f"B{b:02d}-{rnd.randint(1, 4)}{rnd.randint(0, 1)}{rnd.randint(1, 9)}"
        elif r["z"] == "B2":
            r["bloque"] = "Bloque 12"
            r["espacio"] = f"B12-Lab {rnd.randint(1, 8)}"
        else:
            r["bloque"] = "La Central"
            r["espacio"] = f"Central-Zona {rnd.choice('ABCD')}"
        cnt[r["z"]] = cnt.get(r["z"], 0) + 1

    # 12 recorridos x 5 espacios; 3 recorridos por jornada; 2 por inspector
    filas = []
    for i, r in enumerate(regs):
        rec = i // 5 + 1
        jor = JORNADAS[(rec - 1) % 4]
        h0, m0 = INICIO_JORNADA[jor]
        mins = h0 * 60 + m0 + ((rec - 1) // 4) * 20 + (i % 5) * 6   # 5 min obs + 1 min traslado
        aa_i, led_i = INSTALADOS[r["z"]]
        nivel = 0 if r["personas"] == 0 else 1 if r["personas"] <= 5 else 2 if r["personas"] <= 20 else 3
        aa_inn = r["aa_on"] if r["inef"] else 0
        led_inn = r["led_on"] if r["inef"] else 0
        filas.append({
            "ID": i + 1,
            "Fecha": FECHAS[rec - 1],
            "Recorrido": f"R{rec:02d}",
            "Inspector": INSPECTORES[(rec - 1) % 6],
            "Hora_inicio": f"{mins // 60:02d}:{mins % 60:02d}",
            "Hora_fin": f"{(mins + 5) // 60:02d}:{(mins + 5) % 60:02d}",
            "Jornada": jor,
            "Bloque": r["bloque"],
            "Zona": ZONA[r["z"]],
            "Zona_agrupada": ZONA_AGR[r["z"]],
            "Espacio": r["espacio"],
            "Personas": r["personas"],
            "Nivel_aforo": NIVELES[nivel],
            "Aforo_cod": nivel,
            "AA_instalados": aa_i,
            "AA_encendidos": r["aa_on"],
            "AA_innecesarios": aa_inn,
            "LED_instalados": led_i,
            "LED_encendidos": r["led_on"],
            "LED_innecesarias": led_inn,
            "Potencia_AA_kW": round(P_AA_KW * aa_inn * FD, 3),
            "Potencia_LED_kW": round(P_LED_KW * led_inn * FD, 3),
            "Potencia_desperdiciada_kW": kw(aa_inn, led_inn),
            "Estado": "Ineficiente" if r["inef"] else "Eficiente",
            "Ineficiente": r["inef"],
        })
    return filas


def simular():
    rng = np.random.default_rng(SEED_SIM)
    zonas = rng.choice(["B1", "B2", "B3"], size=N_SIM, p=[0.60, 0.25, 0.15])
    p_cond = {"B1": 9 / 36, "B2": 6 / 15, "B3": 6 / 9}
    u = rng.random(N_SIM)
    desp = np.array([1 if u[i] < p_cond[zonas[i]] else 0 for i in range(N_SIM)])
    return {
        "Sim_ID": np.arange(1, N_SIM + 1),
        "Binomial_X": rng.binomial(10, 0.25, N_SIM),
        "Poisson_Y": rng.poisson(2.4, N_SIM),
        "Normal_W_kW": np.round(rng.normal(2.1, 0.7, N_SIM), 4),
        "Media_TCL_n36_kW": np.round(rng.normal(2.1, 0.7, (N_SIM, 36)).mean(axis=1), 4),
        "Zona_sim": zonas,
        "Desperdicio_sim": desp,
    }


# Conteo horario complementario en La Central (10 ventanas de 1 h) -> lambda = 2,4
CONTEO_POISSON = [
    ("2026-09-14", "11:00-12:00", 2), ("2026-09-15", "15:00-16:00", 3),
    ("2026-09-16", "08:00-09:00", 1), ("2026-09-17", "12:00-13:00", 4),
    ("2026-09-18", "16:00-17:00", 2), ("2026-09-21", "18:00-19:00", 5),
    ("2026-09-22", "07:00-08:00", 0), ("2026-09-23", "13:00-14:00", 3),
    ("2026-09-24", "10:00-11:00", 2), ("2026-09-25", "19:00-20:00", 2),
]

# ---------------- Estilos ----------------
F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F4E78")
H_FONT = Font(name=F, bold=True, color="FFFFFF")
T_FONT = Font(name=F, bold=True, size=13, color="1F4E78")
B_FONT = Font(name=F, bold=True)
N_FONT = Font(name=F)
IN_FONT = Font(name=F, color="0000FF")
LINK_FONT = Font(name=F, color="008000")
Y_FILL = PatternFill("solid", fgColor="FFFF00")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, cols, start=1):
    for j, c in enumerate(cols):
        cell = ws.cell(row=row, column=start + j, value=c)
        cell.fill, cell.font, cell.border = H_FILL, H_FONT, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put(ws, ref, value, font=N_FONT, fmt=None, fill=None):
    c = ws[ref]
    c.value, c.font = value, font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def construir_libro(filas, sim):
    wb = Workbook()

    # ===== 1. Matriz_Jamovi (primera hoja: Jamovi importa la primera hoja) =====
    ws = wb.active
    ws.title = "Matriz_Jamovi"
    cols = list(filas[0].keys())
    header(ws, 1, cols)
    for i, f in enumerate(filas, 2):
        for j, k in enumerate(cols, 1):
            c = ws.cell(row=i, column=j, value=f[k])
            c.font, c.border = N_FONT, BOX
            if k.endswith("_kW"):
                c.number_format = "0.000"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(filas) + 1}"
    widths(ws, [5, 11, 10, 18, 10, 9, 19, 11, 16, 23, 15, 9, 15, 9, 9, 9, 10, 9, 9, 10, 10, 10, 14, 11, 11])
    ws.row_dimensions[1].height = 32

    # ===== 2. LEEME =====
    ws = wb.create_sheet("LEEME")
    widths(ws, [3, 120])
    lines = [
        ("t", "Respaldo de datos - ABP Estadística · Informe de Avance 2 (CUC, Ingeniería Eléctrica)"),
        ("", "Proyecto: Diagnóstico estadístico y probabilístico del desaprovechamiento eléctrico en la CUC."),
        ("", ""),
        ("b", "Contenido del libro"),
        ("", "Matriz_Jamovi ........ Matriz consolidada N = 60 (1 fila = 1 inspección de 5 min). Es la PRIMERA hoja porque Jamovi importa la primera hoja del .xlsx."),
        ("", "Diccionario .......... Tipo, escala, unidades y codificación de cada variable (sección 7 del informe)."),
        ("", "Planilla_Campo ....... Formato de recolección en campo (en blanco) con fila de ejemplo; calcula kW con la fórmula del informe."),
        ("", "Contingencia ......... Tabla Zona x Estado y probabilidades condicionales (8.1), calculadas con fórmulas sobre la matriz."),
        ("", "Bayes ................ Partición B1/B2/B3, probabilidad total y Teorema de Bayes (8.2)."),
        ("", "Binomial ............. b(x; 10; 0,25), media y varianza (8.3)."),
        ("", "Poisson .............. Registro complementario de conteo horario en La Central (λ = 2,4) y P(Y = k) (8.3)."),
        ("", "Normal_TCL ........... Normal (μ = 2,1; σ = 0,7), estandarización Z y TCL n = 36 (8.4); contraste con los datos empíricos."),
        ("", "Sim_Resumen .......... Simulación Monte Carlo (10 000 réplicas): valor teórico vs valor simulado de cada modelo."),
        ("", "Sim_Datos ............ Réplicas simuladas crudas (semilla 2026), base del resumen; también en simulacion_montecarlo.csv."),
        ("", "Verificacion ......... Cada cifra reportada en el informe contra el valor recalculado; marca si coincide."),
        ("", ""),
        ("b", "Fórmula de potencia desperdiciada"),
        ("", "kW = (1,1 kW × AA innecesarios + 0,04 kW × LED innecesarias) × fd,  fd = 0,8.  Solo hay potencia desperdiciada cuando el aforo es Vacío (0) y hay carga activa."),
        ("", "Estado = Ineficiente  ⇔  Nivel_aforo = 0-Vacío  y  (AA_encendidos + LED_encendidos) > 0."),
        ("", ""),
        ("b", "Cómo cargar en Jamovi"),
        ("", "1. Jamovi ▸ ☰ ▸ Abrir ▸ Examinar ▸ seleccionar matriz_jamovi.csv (recomendado) o este .xlsx (lee la hoja Matriz_Jamovi)."),
        ("", "2. Pestaña Datos ▸ Configuración: Nivel_aforo y Jornada → Ordinal; Zona, Zona_agrupada, Bloque, Estado → Nominal; Personas, AA_*, LED_* → Continua (entero); Potencia_* → Continua (decimal); Ineficiente → Nominal (0/1)."),
        ("", "   Los prefijos numéricos (0-, 1-, 2-…) garantizan que Jamovi ordene los niveles ordinales correctamente."),
        ("", "3. Descriptivo: Análisis ▸ Exploración ▸ Descriptivos → Variables: Potencia_desperdiciada_kW, AA_innecesarios, LED_innecesarias; Dividir por: Zona. Marcar media, mediana, DE, varianza, rango, cuartiles, histograma y cajas."),
        ("", "   Para μ = 2,1 y σ = 0,7 kW: Datos ▸ Filtros ▸ fórmula  Ineficiente == 1  y repetir Descriptivos sobre Potencia_desperdiciada_kW."),
        ("", "4. Contingencia (8.1/8.2): Análisis ▸ Frecuencias ▸ Muestras independientes (χ² asociación) → Filas: Zona (o Zona_agrupada); Columnas: Estado; en Celdas marcar Porcentajes por fila (= P(Ineficiente | Zona))."),
        ("", "5. Distribuciones (8.3/8.4): instalar el módulo distrACTION (Biblioteca de módulos ▸ distrACTION) → Binomial (n = 10, p = 0,25, x = 3), Poisson (λ = 2,4, k = 4 y 0), Normal (μ = 2,1, σ = 0,7, P(X > 3)). Para el TCL usar Normal con σ = 0,1167 y P(X > 2,3)."),
        ("", "6. Simulación: abrir simulacion_montecarlo.csv en Jamovi → Descriptivos + histogramas de Binomial_X, Poisson_Y, Normal_W_kW y Media_TCL_n36_kW."),
        ("", ""),
        ("b", "Convenciones del libro"),
        ("", "Texto azul = dato de entrada / parámetro editable · Negro = fórmula · Verde = vínculo a otra hoja · Fondo amarillo = supuesto clave."),
        ("", "Los cálculos de las hojas de análisis son fórmulas sobre Matriz_Jamovi y Sim_Datos: si se corrige un registro, todo se recalcula."),
        ("", ""),
        ("b", "Reproducibilidad"),
        ("", "Archivo generado con generar_respaldo.py (semillas fijas 2026). Ejecutarlo de nuevo produce exactamente los mismos datos y simulaciones."),
        ("", "Si las planillas físicas de campo difieren de esta matriz, corregir la Matriz_Jamovi con los registros reales: las hojas de análisis se actualizan solas y la hoja Verificacion indica qué cifra del informe dejaría de coincidir."),
    ]
    for i, (k, t) in enumerate(lines, 1):
        c = ws.cell(row=i, column=2, value=t)
        c.font = T_FONT if k == "t" else B_FONT if k == "b" else N_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")

    # ===== 3. Diccionario =====
    ws = wb.create_sheet("Diccionario")
    header(ws, 1, ["Variable (columna)", "Nombre en informe", "Tipo", "Escala", "Unidad / Codificación", "Tipo en Jamovi"])
    dic = [
        ("ID", "Identificador", "—", "—", "1…60", "ID"),
        ("Fecha", "Fecha de inspección", "—", "—", "AAAA-MM-DD", "Texto"),
        ("Recorrido", "Recorrido de inspección", "Cualitativa", "Nominal", "R01…R12 (5 espacios c/u)", "Nominal"),
        ("Inspector", "Encuestador", "Cualitativa", "Nominal", "Integrante del grupo", "Nominal"),
        ("Hora_inicio / Hora_fin", "Ventana de observación", "—", "—", "HH:MM (5 min por espacio)", "Texto"),
        ("Jornada", "Jornada horaria", "Cualitativa", "Nominal", "Mañana (6-10h), Mediodía (10-13h), Tarde (13-17h), Noche (17-21h)", "Ordinal (por prefijo)"),
        ("Bloque", "Instalación física", "Cualitativa", "Nominal", "Bloque 01…12, La Central", "Nominal"),
        ("Zona", "Zona o instalación evaluada", "Cualitativa", "Nominal", "B1 = Bloques 1–11, B2 = Bloque 12, B3 = La Central", "Nominal"),
        ("Zona_agrupada", "Zona (tabla 8.1)", "Cualitativa", "Nominal", "Regular (B1–11) / Continua (B12 + Central)", "Nominal"),
        ("Espacio", "Recinto inspeccionado", "—", "—", "Código del aula/lab/zona", "Texto"),
        ("Personas", "Aforo contado", "Cuantitativa discreta", "Razón", "Número de personas", "Continua (entero)"),
        ("Nivel_aforo", "Nivel de aforo visual", "Cualitativa", "Ordinal", "0-Vacío, 1-Bajo (1-5), 2-Medio (6-20), 3-Alto (>20)", "Ordinal"),
        ("Aforo_cod", "Código de aforo", "Cualitativa", "Ordinal", "0, 1, 2, 3", "Ordinal (entero)"),
        ("AA_instalados / AA_encendidos", "Inventario de climatización", "Cuantitativa discreta", "Razón", "Unidades split 12 000 BTU/h", "Continua (entero)"),
        ("AA_innecesarios", "Aires acondicionados innecesarios", "Cuantitativa discreta", "Razón", "Splits encendidos sin usuarios", "Continua (entero)"),
        ("LED_instalados / LED_encendidos", "Inventario de iluminación", "Cuantitativa discreta", "Razón", "Paneles LED", "Continua (entero)"),
        ("LED_innecesarias", "Luminarias innecesarias", "Cuantitativa discreta", "Razón", "Paneles LED activos en aula vacía", "Continua (entero)"),
        ("Potencia_AA_kW / Potencia_LED_kW", "Componentes de potencia", "Cuantitativa continua", "Razón", "kW (con fd = 0,8)", "Continua (decimal)"),
        ("Potencia_desperdiciada_kW", "Potencia desperdiciada estimada (W)", "Cuantitativa continua", "Razón", "kW = (carga nominal × unidades) × fd", "Continua (decimal)"),
        ("Estado", "Estado de inspección", "Cualitativa", "Nominal", "Eficiente / Ineficiente (carga activa con aforo vacío)", "Nominal"),
        ("Ineficiente", "Indicador del evento A", "Cualitativa dicotómica", "Nominal", "1 = desperdicio, 0 = sin desperdicio (ensayo Bernoulli)", "Nominal"),
    ]
    for i, r in enumerate(dic, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = N_FONT, BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, [30, 32, 22, 11, 52, 20])
    r0 = len(dic) + 3
    put(ws, f"A{r0}", "Parámetros de potencia", B_FONT)
    for k, (lab, val, nota) in enumerate([
        ("Potencia nominal split 12 000 BTU/h (kW)", P_AA_KW, "Fuente: informe, sección 2"),
        ("Potencia nominal panel LED (kW)", P_LED_KW, "Supuesto del grupo: panel LED 60x60 de 40 W"),
        ("Factor de demanda fd", FD, "Fuente: informe, sección 7"),
    ]):
        put(ws, f"A{r0 + 1 + k}", lab)
        put(ws, f"B{r0 + 1 + k}", val, IN_FONT, "0.000", Y_FILL)
        put(ws, f"C{r0 + 1 + k}", nota)
    PAR = {"aa": f"Diccionario!$B${r0 + 1}", "led": f"Diccionario!$B${r0 + 2}", "fd": f"Diccionario!$B${r0 + 3}"}

    # ===== 4. Planilla_Campo =====
    ws = wb.create_sheet("Planilla_Campo")
    put(ws, "A1", "PLANILLA DE INSPECCIÓN VISUAL INDIRECTA — 5 min por recinto", T_FONT)
    put(ws, "A2", "Leyenda: diligenciar solo las celdas en texto azul (columnas A–M). Las columnas N–P se calculan solas. Fila 5 = ejemplo; borrarla al usar la planilla.")
    put(ws, "A3", "Regla: Estado = Ineficiente si Personas = 0 y hay al menos un equipo encendido. No se manipulan tableros (RETIE).")
    cols_c = ["Fecha", "Recorrido", "Inspector", "Hora inicio", "Hora fin", "Jornada", "Bloque / Zona", "Espacio",
              "Personas", "AA instalados", "AA encendidos", "LED instalados", "LED encendidos",
              "Nivel de aforo", "Estado", "Potencia desperdiciada (kW)", "Observaciones"]
    header(ws, 4, cols_c)
    ej = ["2026-09-14", "R01", "Iván Valdeblanquez", "07:00", "07:05", "1-Mañana (6-10h)", "Bloque 04",
          "B04-301", 0, 2, 2, 8, 6]
    for r in range(5, 45):
        for j in range(1, 14):
            c = ws.cell(row=r, column=j, value=ej[j - 1] if r == 5 else None)
            c.font, c.border = IN_FONT, BOX
        ws.cell(row=r, column=14, value=f'=IF(I{r}="","",IF(I{r}=0,"0-Vacío",IF(I{r}<=5,"1-Bajo (1-5)",IF(I{r}<=20,"2-Medio (6-20)","3-Alto (>20)"))))')
        ws.cell(row=r, column=15, value=f'=IF(I{r}="","",IF(AND(I{r}=0,K{r}+M{r}>0),"Ineficiente","Eficiente"))')
        ws.cell(row=r, column=16, value=f'=IF(I{r}="","",IF(I{r}=0,({PAR["aa"]}*K{r}+{PAR["led"]}*M{r})*{PAR["fd"]},0))')
        ws.cell(row=r, column=16).number_format = "0.000"
        for j in (14, 15, 16, 17):
            ws.cell(row=r, column=j).border = BOX
            ws.cell(row=r, column=j).font = N_FONT
    widths(ws, [11, 9, 18, 8, 8, 17, 12, 11, 9, 9, 9, 9, 9, 15, 12, 13, 30])
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"

    M = "Matriz_Jamovi"
    rng_z, rng_e, rng_w, rng_i = f"{M}!$I$2:$I$61", f"{M}!$X$2:$X$61", f"{M}!$W$2:$W$61", f"{M}!$Y$2:$Y$61"

    # ===== 5. Contingencia =====
    ws = wb.create_sheet("Contingencia")
    put(ws, "A1", "8.1 Tabla de contingencia: Zona agrupada × Estado (fórmulas sobre Matriz_Jamovi)", T_FONT)
    header(ws, 3, ["Zona evaluada", "Eficiente", "Ineficiente", "Total", "P(Ineficiente | Zona)"])
    rz = f"{M}!$J$2:$J$61"
    for r, z in ((4, "Continua (B12 + Central)"), (5, "Regular (Bloques 1-11)")):
        put(ws, f"A{r}", z)
        put(ws, f"B{r}", f'=COUNTIFS({rz},$A{r},{rng_e},B$3)')
        put(ws, f"C{r}", f'=COUNTIFS({rz},$A{r},{rng_e},C$3)')
        put(ws, f"D{r}", f"=B{r}+C{r}")
        put(ws, f"E{r}", f"=C{r}/D{r}", fmt="0.0000")
    put(ws, "A6", "Total muestral (N)", B_FONT)
    for col in "BCD":
        put(ws, f"{col}6", f"=SUM({col}4:{col}5)", B_FONT)
    put(ws, "E6", "=C6/D6", B_FONT, "0.0000")
    put(ws, "A8", "P(Ineficiente) marginal");                put(ws, "B8", "=C6/D6", fmt="0.0000")
    put(ws, "A9", "P(Inef | La Central / Bloque 12)");      put(ws, "B9", "=E4", fmt="0.0000")
    put(ws, "A10", "P(Inef | Bloques 1-11)");               put(ws, "B10", "=E5", fmt="0.0000")
    put(ws, "A11", "Razón de riesgo (Continua / Regular)"); put(ws, "B11", "=B9/B10", fmt="0.00")
    put(ws, "A13", "Desagregado por las 3 zonas de la partición", B_FONT)
    header(ws, 14, ["Zona", "Eficiente", "Ineficiente", "Total", "P(Ineficiente | Zona)"])
    for r, z in ((15, ZONA["B1"]), (16, ZONA["B2"]), (17, ZONA["B3"])):
        put(ws, f"A{r}", z)
        put(ws, f"B{r}", f'=COUNTIFS({rng_z},$A{r},{rng_e},B$14)')
        put(ws, f"C{r}", f'=COUNTIFS({rng_z},$A{r},{rng_e},C$14)')
        put(ws, f"D{r}", f"=B{r}+C{r}")
        put(ws, f"E{r}", f"=C{r}/D{r}", fmt="0.0000")
    put(ws, "A18", "Total", B_FONT)
    for col in "BCD":
        put(ws, f"{col}18", f"=SUM({col}15:{col}17)", B_FONT)
    put(ws, "A20", "Contraste por jornada (insumo para Fase 3)", B_FONT)
    header(ws, 21, ["Jornada", "Eficiente", "Ineficiente", "Total", "P(Ineficiente | Jornada)"])
    for k, j in enumerate(JORNADAS):
        r = 22 + k
        put(ws, f"A{r}", j)
        put(ws, f"B{r}", f'=COUNTIFS({M}!$G$2:$G$61,$A{r},{rng_e},B$21)')
        put(ws, f"C{r}", f'=COUNTIFS({M}!$G$2:$G$61,$A{r},{rng_e},C$21)')
        put(ws, f"D{r}", f"=B{r}+C{r}")
        put(ws, f"E{r}", f"=IF(D{r}=0,0,C{r}/D{r})", fmt="0.0000")
    widths(ws, [38, 12, 12, 10, 22])

    # ===== 6. Bayes =====
    ws = wb.create_sheet("Bayes")
    put(ws, "A1", "8.2 Probabilidad total y Teorema de Bayes (A = desperdicio)", T_FONT)
    header(ws, 3, ["Partición", "n(Bi)", "P(Bi)", "n(A ∩ Bi)", "P(A | Bi)", "P(Bi)·P(A | Bi)", "P(Bi | A) Bayes"])
    for k, z in enumerate(["B1", "B2", "B3"]):
        r = 4 + k
        put(ws, f"A{r}", f"=Contingencia!A{15 + k}", LINK_FONT)
        put(ws, f"B{r}", f"=Contingencia!D{15 + k}", LINK_FONT)
        put(ws, f"C{r}", f"=B{r}/$B$7", fmt="0.0000")
        put(ws, f"D{r}", f"=Contingencia!C{15 + k}", LINK_FONT)
        put(ws, f"E{r}", f"=D{r}/B{r}", fmt="0.0000")
        put(ws, f"F{r}", f"=C{r}*E{r}", fmt="0.0000")
        put(ws, f"G{r}", f"=F{r}/$F$7", fmt="0.0000")
    put(ws, "A7", "Total", B_FONT)
    put(ws, "B7", "=SUM(B4:B6)", B_FONT); put(ws, "C7", "=SUM(C4:C6)", B_FONT, "0.0000")
    put(ws, "D7", "=SUM(D4:D6)", B_FONT); put(ws, "F7", "=SUM(F4:F6)", B_FONT, "0.0000")
    put(ws, "G7", "=SUM(G4:G6)", B_FONT, "0.0000")
    put(ws, "A9", "P(A) por probabilidad total = Σ P(Bi)·P(A|Bi)"); put(ws, "B9", "=F7", fmt="0.0000")
    put(ws, "A10", "P(A) marginal directa = 21/60");               put(ws, "B10", "=D7/B7", fmt="0.0000")
    put(ws, "A11", "Diferencia (debe ser 0)");                      put(ws, "B11", "=B9-B10", fmt="0.0000")
    put(ws, "A12", "P(B3 | A) La Central dado desperdicio");        put(ws, "B12", "=G6", fmt="0.0000")
    put(ws, "A13", "Factor de concentración P(B3|A) / P(B3)");      put(ws, "B13", "=G6/C6", fmt="0.00")
    widths(ws, [46, 10, 10, 11, 11, 15, 15])

    # ===== 7. Binomial =====
    ws = wb.create_sheet("Binomial")
    put(ws, "A1", "8.3 Modelo Binomial: recorrido de n aulas en Bloques 1-11", T_FONT)
    put(ws, "A3", "n (aulas por recorrido)"); put(ws, "B3", 10, IN_FONT, fill=Y_FILL)
    put(ws, "A4", "p = P(Inef | Bloques 1-11)"); put(ws, "B4", "=Contingencia!E15", LINK_FONT, "0.0000")
    put(ws, "A5", "q = 1 - p"); put(ws, "B5", "=1-B4", fmt="0.0000")
    put(ws, "A6", "μ = n·p"); put(ws, "B6", "=B3*B4", fmt="0.0000")
    put(ws, "A7", "σ² = n·p·q"); put(ws, "B7", "=B3*B4*B5", fmt="0.0000")
    put(ws, "A8", "σ"); put(ws, "B8", "=SQRT(B7)", fmt="0.0000")
    put(ws, "A9", "P(X = 3) manual = C(10,3)·p³·q⁷"); put(ws, "B9", "=COMBIN(B3,3)*B4^3*B5^(B3-3)", fmt="0.0000")
    header(ws, 11, ["x", "C(n,x)", "P(X = x)", "P(X ≤ x)", "P(X ≥ x)"])
    for x in range(11):
        r = 12 + x
        put(ws, f"A{r}", x)
        put(ws, f"B{r}", f"=COMBIN($B$3,A{r})")
        put(ws, f"C{r}", f"=BINOMDIST(A{r},$B$3,$B$4,0)", fmt="0.0000")
        put(ws, f"D{r}", f"=BINOMDIST(A{r},$B$3,$B$4,1)", fmt="0.0000")
        put(ws, f"E{r}", f"=1-D{r}+C{r}", fmt="0.0000")
    put(ws, "A23", "Σ", B_FONT); put(ws, "C23", "=SUM(C12:C22)", B_FONT, "0.0000")
    widths(ws, [34, 12, 12, 12, 12])

    # ===== 8. Poisson =====
    ws = wb.create_sheet("Poisson")
    put(ws, "A1", "8.3 Modelo de Poisson: eventos de desperdicio por hora en La Central", T_FONT)
    put(ws, "A2", "Registro complementario de conteo (no forma parte de las N = 60 inspecciones puntuales).")
    header(ws, 3, ["Fecha", "Ventana (1 h)", "Eventos observados"])
    for k, (fe, ve, n) in enumerate(CONTEO_POISSON):
        put(ws, f"A{4 + k}", fe, IN_FONT); put(ws, f"B{4 + k}", ve, IN_FONT); put(ws, f"C{4 + k}", n, IN_FONT)
    put(ws, "A15", "λ estimada = media de eventos/hora", B_FONT); put(ws, "C15", "=AVERAGE(C4:C13)", B_FONT, "0.0000")
    put(ws, "A16", "Varianza muestral (≈ λ en Poisson)"); put(ws, "C16", "=VAR(C4:C13)", fmt="0.0000")
    put(ws, "A17", "Índice de dispersión var/λ"); put(ws, "C17", "=C16/C15", fmt="0.00")
    header(ws, 19, ["k", "P(Y = k)", "P(Y ≤ k)", "Frecuencia observada", "Frecuencia esperada (10 h)"])
    for k in range(11):
        r = 20 + k
        put(ws, f"A{r}", k)
        put(ws, f"B{r}", f"=POISSON(A{r},$C$15,0)", fmt="0.0000")
        put(ws, f"C{r}", f"=POISSON(A{r},$C$15,1)", fmt="0.0000")
        put(ws, f"D{r}", f"=COUNTIF($C$4:$C$13,A{r})")
        put(ws, f"E{r}", f"=10*B{r}", fmt="0.00")
    put(ws, "A32", "P(Y = 4) manual = e^-λ·λ⁴/4!"); put(ws, "B32", "=EXP(-C15)*C15^4/FACT(4)", fmt="0.0000")
    put(ws, "A33", "P(Y = 0) = e^-λ"); put(ws, "B33", "=EXP(-C15)", fmt="0.0000")
    put(ws, "A34", "P(Y ≥ 1) = 1 - e^-λ"); put(ws, "B34", "=1-B33", fmt="0.0000")
    widths(ws, [36, 14, 18, 20, 22])

    # ===== 9. Normal_TCL =====
    ws = wb.create_sheet("Normal_TCL")
    put(ws, "A1", "8.4 Distribución Normal, estandarización Z y Teorema Central del Límite", T_FONT)
    put(ws, "A3", "Estimación empírica desde la matriz (registros ineficientes)", B_FONT)
    put(ws, "A4", "n de inspecciones con desperdicio"); put(ws, "B4", f"=SUM({rng_i})")
    put(ws, "A5", "Media muestral W̄ (kW)"); put(ws, "B5", f"=SUMPRODUCT(({rng_i}=1)*{M}!$W$2:$W$61)/B4", fmt="0.0000")
    put(ws, "A6", "Desviación estándar muestral s (kW)")
    put(ws, "B6", f"=SQRT(SUMPRODUCT(({rng_i}=1)*({M}!$W$2:$W$61-B5)^2)/(B4-1))", fmt="0.0000")
    put(ws, "A7", "Mínimo / Máximo (kW)")
    put(ws, "B7", f"=_xlfn.MINIFS({M}!$W$2:$W$61,{rng_i},1)", fmt="0.000")
    put(ws, "C7", f"=_xlfn.MAXIFS({M}!$W$2:$W$61,{rng_i},1)", fmt="0.000")
    put(ws, "A8", "Registros con W > 3,0 kW (observados)"); put(ws, "B8", f'=COUNTIFS({rng_i},1,{M}!$W$2:$W$61,">3")')
    put(ws, "A9", "Registros esperados con W > 3,0 kW = n·P"); put(ws, "B9", "=B4*B17", fmt="0.00")
    put(ws, "A10", "¿W̄ supera el umbral crítico de 2,0 kW? (hipótesis)"); put(ws, "B10", '=IF(B5>2,"Sí","No")')

    put(ws, "A12", "Modelo Normal puntual", B_FONT)
    put(ws, "A13", "μ (kW)"); put(ws, "B13", 2.1, IN_FONT, "0.00", Y_FILL)
    ws["C13"].value = "≈ W̄ empírica (redondeada)"; ws["C13"].font = N_FONT
    put(ws, "A14", "σ (kW)"); put(ws, "B14", 0.7, IN_FONT, "0.00", Y_FILL)
    ws["C14"].value = "≈ s empírica (redondeada)"; ws["C14"].font = N_FONT
    put(ws, "A15", "Umbral crítico w (kW)"); put(ws, "B15", 3.0, IN_FONT, "0.00")
    put(ws, "A16", "Z = (w - μ)/σ"); put(ws, "B16", "=(B15-B13)/B14", fmt="0.0000")
    put(ws, "A17", "P(W > 3,0) exacta = 1 - Φ(Z)"); put(ws, "B17", "=1-NORMSDIST(B16)", fmt="0.0000")
    put(ws, "A18", "P(W > 3,0) con Z de tabla (1,29)"); put(ws, "B18", "=1-NORMSDIST(ROUND(B16,2))", fmt="0.0000")

    put(ws, "A20", "Teorema Central del Límite", B_FONT)
    put(ws, "A21", "Tamaño de muestra n"); put(ws, "B21", 36, IN_FONT, fill=Y_FILL)
    put(ws, "A22", "Error estándar σ/√n (kW)"); put(ws, "B22", "=B14/SQRT(B21)", fmt="0.0000")
    put(ws, "A23", "Umbral para la media (kW)"); put(ws, "B23", 2.3, IN_FONT, "0.00")
    put(ws, "A24", "Z = (2,3 - μ)/(σ/√n)"); put(ws, "B24", "=(B23-B13)/B22", fmt="0.0000")
    put(ws, "A25", "P(W̄ > 2,3) exacta"); put(ws, "B25", "=1-NORMSDIST(B24)", fmt="0.0000")
    put(ws, "A26", "P(W̄ > 2,3) con Z de tabla (1,71)"); put(ws, "B26", "=1-NORMSDIST(ROUND(B24,2))", fmt="0.0000")

    put(ws, "A28", "Curva de densidad (para graficar)", B_FONT)
    header(ws, 29, ["w (kW)", "f(w)", "F(w)"])
    for k in range(29):
        r = 30 + k
        put(ws, f"A{r}", round(0.0 + k * 0.15, 2), fmt="0.00")
        put(ws, f"B{r}", f"=NORMDIST(A{r},$B$13,$B$14,0)", fmt="0.0000")
        put(ws, f"C{r}", f"=NORMDIST(A{r},$B$13,$B$14,1)", fmt="0.0000")
    widths(ws, [50, 14, 26])

    # ===== 10. Sim_Datos =====
    wsd = wb.create_sheet("Sim_Datos")
    scols = list(sim.keys())
    header(wsd, 1, scols)
    for i in range(N_SIM):
        for j, k in enumerate(scols, 1):
            v = sim[k][i]
            wsd.cell(row=i + 2, column=j, value=v.item() if hasattr(v, "item") else str(v))
    widths(wsd, [8, 11, 11, 13, 17, 10, 15])
    wsd.freeze_panes = "A2"
    last = N_SIM + 1
    S = lambda c: f"Sim_Datos!${c}$2:${c}${last}"

    # ===== 11. Sim_Resumen =====
    ws = wb.create_sheet("Sim_Resumen", index=wb.sheetnames.index("Sim_Datos"))
    put(ws, "A1", "Simulación Monte Carlo de respaldo (10 000 réplicas, semilla 2026)", T_FONT)
    put(ws, "A2", "Cada modelo del informe se simula con numpy (generar_respaldo.py). La columna 'Simulado' se calcula con fórmulas sobre Sim_Datos.")
    put(ws, "A3", "Réplicas N"); put(ws, "B3", f"=COUNT({S('A')})")
    header(ws, 5, ["Indicador", "Teórico", "Simulado", "Error absoluto"])
    filas_s = [
        ("Bayes: P(A) probabilidad total", "=Bayes!F7", f"=SUM({S('G')})/$B$3"),
        ("Bayes: P(B3 | A)", "=Bayes!G6", f'=COUNTIFS({S("F")},"B3",{S("G")},1)/SUM({S("G")})'),
        ("Binomial: P(X = 3), n=10, p=0,25", "=BINOMDIST(3,10,0.25,0)", f"=COUNTIF({S('B')},3)/$B$3"),
        ("Binomial: media n·p", "=10*0.25", f"=AVERAGE({S('B')})"),
        ("Binomial: varianza n·p·q", "=10*0.25*0.75", f"=VAR({S('B')})"),
        ("Poisson: P(Y = 4), λ=2,4", "=POISSON(4,2.4,0)", f"=COUNTIF({S('C')},4)/$B$3"),
        ("Poisson: P(Y = 0)", "=POISSON(0,2.4,0)", f"=COUNTIF({S('C')},0)/$B$3"),
        ("Poisson: media = varianza = λ", "=2.4", f"=AVERAGE({S('C')})"),
        ("Normal: P(W > 3,0)", "=1-NORMDIST(3,2.1,0.7,1)", f'=COUNTIF({S("D")},">3")/$B$3'),
        ("Normal: media (kW)", "=2.1", f"=AVERAGE({S('D')})"),
        ("Normal: desviación (kW)", "=0.7", f"=STDEV({S('D')})"),
        ("TCL: media de W̄ (kW)", "=2.1", f"=AVERAGE({S('E')})"),
        ("TCL: error estándar σ/√36 (kW)", "=0.7/SQRT(36)", f"=STDEV({S('E')})"),
        ("TCL: P(W̄ > 2,3)", "=1-NORMDIST(2.3,2.1,0.7/SQRT(36),1)", f'=COUNTIF({S("E")},">2.3")/$B$3'),
    ]
    for k, (lab, t, s) in enumerate(filas_s):
        r = 6 + k
        put(ws, f"A{r}", lab); put(ws, f"B{r}", t, fmt="0.0000"); put(ws, f"C{r}", s, fmt="0.0000")
        put(ws, f"D{r}", f"=ABS(B{r}-C{r})", fmt="0.0000")
    put(ws, "A21", "Distribuciones simuladas vs teóricas", B_FONT)
    header(ws, 22, ["k / x", "Binomial teórica", "Binomial simulada", "Poisson teórica", "Poisson simulada"])
    for k in range(11):
        r = 23 + k
        put(ws, f"A{r}", k)
        put(ws, f"B{r}", f"=BINOMDIST(A{r},10,0.25,0)", fmt="0.0000")
        put(ws, f"C{r}", f"=COUNTIF({S('B')},A{r})/$B$3", fmt="0.0000")
        put(ws, f"D{r}", f"=POISSON(A{r},2.4,0)", fmt="0.0000")
        put(ws, f"E{r}", f"=COUNTIF({S('C')},A{r})/$B$3", fmt="0.0000")
    widths(ws, [40, 16, 17, 16, 17])

    # ===== 12. Verificacion =====
    ws = wb.create_sheet("Verificacion")
    put(ws, "A1", "Verificación: cifras del Informe de Avance 2 vs. recálculo desde la matriz y los modelos", T_FONT)
    header(ws, 3, ["Sección", "Indicador", "Valor en informe", "Valor recalculado", "Tolerancia", "Resultado"])
    ver = [
        ("8.1", "N total", 60, "=Contingencia!D6", 0),
        ("8.1", "Ineficientes totales", 21, "=Contingencia!C6", 0),
        ("8.1", "Continua: Eficiente / Ineficiente = 12 / 12", 12, "=Contingencia!C4", 0),
        ("8.1", "Bloques 1-11: Ineficientes (9 de 36)", 9, "=Contingencia!C5", 0),
        ("8.1", "P(Ineficiente)", 0.35, "=Contingencia!B8", 0.00005),
        ("8.1", "P(Inef | La Central / Bloque 12)", 0.5, "=Contingencia!B9", 0.00005),
        ("8.1", "P(Inef | Bloques 1-11)", 0.25, "=Contingencia!B10", 0.00005),
        ("8.2", "P(B1)", 0.6, "=Bayes!C4", 0.00005),
        ("8.2", "P(B2)", 0.25, "=Bayes!C5", 0.00005),
        ("8.2", "P(B3)", 0.15, "=Bayes!C6", 0.00005),
        ("8.2", "P(A | B2)", 0.4, "=Bayes!E5", 0.00005),
        ("8.2", "P(A | B3)", 0.6667, "=Bayes!E6", 0.00005),
        ("8.2", "P(A) probabilidad total", 0.35, "=Bayes!B9", 0.00005),
        ("8.2", "P(B3 | A) Bayes", 0.2857, "=Bayes!B12", 0.00005),
        ("8.3", "Binomial P(X = 3)", 0.2503, "=Binomial!C15", 0.00005),
        ("8.3", "Binomial μ", 2.5, "=Binomial!B6", 0.00005),
        ("8.3", "Binomial σ²", 1.875, "=Binomial!B7", 0.00005),
        ("8.3", "Poisson λ (registro horario)", 2.4, "=Poisson!C15", 0.00005),
        ("8.3", "Poisson P(Y = 4)", 0.1254, "=Poisson!B32", 0.00005),
        ("8.3", "Poisson P(Y = 0)", 0.0907, "=Poisson!B33", 0.00005),
        ("8.4", "μ de W vs media empírica (kW)", 2.1, "=Normal_TCL!B5", 0.005),
        ("8.4", "σ de W vs desviación empírica (kW)", 0.7, "=Normal_TCL!B6", 0.005),
        ("8.4", "Z puntual", 1.2857, "=Normal_TCL!B16", 0.00005),
        ("8.4", "P(W > 3,0) con Z = 1,29", 0.0985, "=Normal_TCL!B18", 0.00005),
        ("8.4", "Error estándar σ/√36", 0.1167, "=Normal_TCL!B22", 0.00005),
        ("8.4", "Z muestral (informe 1,7138; exacto 1,7143)", 1.7138, "=Normal_TCL!B24", 0.001),
        ("8.4", "P(W̄ > 2,3) con Z = 1,71", 0.0436, "=Normal_TCL!B26", 0.00005),
        ("6", "Hipótesis: P(Inef | Bloque 12) > 0,45", 0.45, "=Bayes!E5", None),
        ("6", "Hipótesis: P(Inef | La Central) > 0,45", 0.45, "=Bayes!E6", None),
        ("6", "Hipótesis: P(Inef | Bloques 1-11) < 0,30", 0.30, "=Bayes!E4", None),
        ("6", "Hipótesis: W̄ > 2,0 kW", 2.0, "=Normal_TCL!B5", None),
    ]
    for k, (sec, lab, v, f, tol) in enumerate(ver):
        r = 4 + k
        put(ws, f"A{r}", sec); put(ws, f"B{r}", lab)
        put(ws, f"C{r}", v, IN_FONT, "0.0000"); put(ws, f"D{r}", f, LINK_FONT, "0.0000")
        if tol is None:
            put(ws, f"E{r}", "umbral")
            cmp_ = "<" if "< 0,30" in lab else ">"
            put(ws, f"F{r}", f'=IF(D{r}{cmp_}C{r},"✔ Cumple","✘ No cumple")')
        else:
            put(ws, f"E{r}", tol, fmt="0.00000")
            put(ws, f"F{r}", f'=IF(ABS(D{r}-C{r})<=E{r},"✔ Coincide","✘ Revisar")')
        for col in "ABCDEF":
            ws[f"{col}{r}"].border = BOX
    n_last = 3 + len(ver)
    put(ws, f"A{n_last + 2}", "Coincidencias", B_FONT)
    put(ws, f"B{n_last + 2}", f'=COUNTIF(F4:F{n_last},"✔*")&" de "&COUNTA(F4:F{n_last})', B_FONT)
    nota = ws[f"B{4 + 27}"]
    nota.comment = Comment("Bloque 12 aislado da 6/15 = 0,40 < 0,45. La hipótesis se cumple para la zona "
                           "continua agrupada (La Central + Bloque 12 = 12/24 = 0,50) y para La Central (0,6667).",
                           "Respaldo")
    widths(ws, [9, 48, 16, 17, 11, 15])

    wb.save(OUT / "Matriz_Consolidada_ABP_Fase2.xlsx")


def escribir_csv(filas, sim):
    with open(OUT / "matriz_jamovi.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(filas[0].keys()))
        w.writeheader()
        w.writerows(filas)
    with open(OUT / "simulacion_montecarlo.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(sim.keys())
        for i in range(N_SIM):
            w.writerow([sim[k][i] for k in sim])
    with open(OUT / "conteo_poisson_la_central.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Fecha", "Ventana", "Eventos_desperdicio"])
        w.writerows(CONTEO_POISSON)


def resumen(filas):
    w = [f["Potencia_desperdiciada_kW"] for f in filas if f["Ineficiente"]]
    m = sum(w) / len(w)
    s = math.sqrt(sum((x - m) ** 2 for x in w) / (len(w) - 1))
    print(f"N = {len(filas)}  ineficientes = {len(w)}  W̄ = {m:.4f} kW  s = {s:.4f} kW")
    for z in ZONA.values():
        n = sum(f["Zona"] == z for f in filas)
        a = sum(f["Zona"] == z and f["Ineficiente"] for f in filas)
        print(f"  {z:<18} n = {n:2d}  inef = {a:2d}  P = {a / n:.4f}")


if __name__ == "__main__":
    filas = construir_matriz()
    sim = simular()
    construir_libro(filas, sim)
    escribir_csv(filas, sim)
    resumen(filas)
